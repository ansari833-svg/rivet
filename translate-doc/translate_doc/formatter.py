"""Format mode: turn a raw ``_translated.docx`` into a clean, footnoted document.

The Translate-mode output is legible line-for-line but reads as walls of
shredded text: footnotes sit inline as plain text, the body is one short
paragraph per source line, and the file is cluttered with ``[[page N]]``
markers, ``[†see log]`` flags and OCR garbage from cover pages.

Format mode parses that structure, promotes the inline footnotes to real Word
footnotes (via Pandoc), reflows the body into prose, and — when an API key is
present — asks Claude to insert paragraph breaks only where the argument
genuinely shifts. The result is ``{name}_formatted.docx``.

Pandoc is a required external dependency for this mode (python-docx cannot
create real Word footnotes).
"""

from __future__ import annotations

import json
import re
import shutil
import statistics
import subprocess
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

MODEL = "claude-sonnet-5"

# Pages with fewer Latin letters than this are treated as un-OCR-able cover /
# title pages and dropped.
LATIN_LETTER_THRESHOLD = 40

# Target chunk size (in words) for the intelligent-paragraphing API calls.
WORDS_PER_CHUNK = 2500

PAGE_MARKER_RE = re.compile(r"\[\[\s*page\s+(\d+)\s*\]\]", re.IGNORECASE)
FOOTNOTE_DEF_RE = re.compile(r"^\((\d+)\)")
SEE_LOG_RE = re.compile(r"\[\s*†?\s*see\s*log\s*\]", re.IGNORECASE)
LEAKED_TAG_RE = re.compile(r"</?(?:sup|sub|u)>", re.IGNORECASE)
TOKEN_RE = re.compile(r"(<<F\d+>>|<<P\d+>>|<<A>>)")
QURAN_RE = re.compile(r"\{([^{}]*)\}")
_SENTENCE_END = ".?!…:؟"

# OCR misreads: CJK symbols/ideographs, kana, Hangul, and the 《》 brackets.
# (U+3040–U+30FF and 《》 fall inside U+3000–U+9FFF but are listed for clarity.)
CJK_RE = re.compile(r"[　-鿿぀-ヿ가-퟿《》]")

# Running-header / title / journal keywords. A short body line matching one of
# these is dropped as a running header rather than kept as content.
HEADER_KEYWORDS = (
    "criterion of prohibited resemblance",
    "al-juday",
    "al-judayʿ",
    "academic journal",
    "journal of",
    "issue no",
    "muḥarram",
    "muharram",
    "october 2016",
)
HEADER_MONTH_YEAR_RE = re.compile(
    r"\b(january|february|march|april|may|june|july|august|september|october|"
    r"november|december|mu[ḥh]arram|safar|rabi[ʿ'`]?|jumada|rajab|sha[ʿ'`]?ban|"
    r"rama[ḍd]an|shawwal|dhu[ʾ'`]?l)\b.*\b\d{4}\b",
    re.IGNORECASE,
)
# A line that is nothing but a page number (optionally bracketed / dashed).
BARE_PAGE_NUM_RE = re.compile(r"^[\[\(\-–—\s]*\d{1,4}[\]\)\-–—\s]*$")

# Discourse markers that should start a new paragraph.
DISCOURSE_RE = re.compile(
    r"^(?:as for\b|the first\b|the second\b|the third\b|then\b|accordingly\b|"
    r"therefore\b|thus\b|in summary\b|chapter\b|section\b|question\b|"
    r"know that\b|he said:|it was narrated\b|on the authority\b|and allah\b|"
    r"allah(?: the exalted| the almighty| almighty)? said)",
    re.IGNORECASE,
)

# Sentence splitting: break after .!?؟… when followed by whitespace then an
# uppercase letter or opening quote (keeps "no. 1" and decimals intact).
SENTENCE_SPLIT_RE = re.compile(r"(?<=[.!?؟…])\s+(?=[\"“A-Z])")
# Sentence punctuation glued to the next word: insert a space.
GLUED_PUNCT_RE = re.compile(r"([.!?؟])([\"“A-Z])")

# Mechanical-paragraphing thresholds (the fallback when the LLM step is off).
_PARA_MARKER_MIN_CHARS = 160   # break BEFORE a discourse marker past this
_PARA_LEN_MIN_SENTENCES = 2    # break AFTER a boundary once this many sentences
_PARA_LEN_MIN_CHARS = 380      # ...and this many chars have accumulated
_PARA_MAX_CHARS = 1100         # last-resort split above this
_PARA_SPLIT_MIN_CHARS = 600    # ...at a ; or : clause end past this

PARAGRAPHING_SYSTEM_PROMPT = """\
You insert paragraph breaks into an already-translated passage. You are given \
running text that has had its paragraphing flattened.

Your ONLY permitted action is to insert paragraph breaks (blank lines) between \
sentences, and only where the argument genuinely shifts — a new point, a new \
hadith, a move from a ruling to its evidence, or the start of a new section.

Absolute rules:
1. Do NOT change, add, remove, reorder, correct, or translate any words, \
punctuation, or characters. The text between breaks must be byte-for-byte the \
same as the input (aside from the whitespace you convert into paragraph breaks).
2. Preserve every placeholder token exactly and in place. Tokens look like \
<<F12>>, <<P5>>, and <<A>>. Never alter, drop, duplicate, translate, or move \
them relative to the surrounding words.
3. Only break BETWEEN sentences, never inside one. If unsure, do not break.
4. Return ONLY the re-paragraphed text — no commentary, no headings, no \
markup other than the blank lines separating paragraphs.
"""


# ---------------------------------------------------------------------------
# Data model
# ---------------------------------------------------------------------------


@dataclass
class RunInfo:
    text: str
    superscript: bool
    highlighted: bool


@dataclass
class ParaInfo:
    text: str
    runs: list[RunInfo]


@dataclass
class Footnote:
    page_num: int
    local_num: int
    text: str
    global_id: int = 0


@dataclass
class Page:
    num: int
    raw_paras: list[ParaInfo] = field(default_factory=list)
    body: list[ParaInfo] = field(default_factory=list)
    footnotes: list[Footnote] = field(default_factory=list)
    skipped: bool = False
    reflowed: str = ""


# ---------------------------------------------------------------------------
# Step 1 — parse structure
# ---------------------------------------------------------------------------


def _para_info(paragraph) -> ParaInfo:
    runs = [
        RunInfo(
            text=r.text,
            superscript=bool(r.font.superscript),
            highlighted=r.font.highlight_color is not None,
        )
        for r in paragraph.runs
    ]
    return ParaInfo(text=paragraph.text, runs=runs)


def _first_nonempty_superscript(p: ParaInfo) -> bool:
    for r in p.runs:
        if r.text.strip():
            return r.superscript
    return False


def _is_footnote_def(p: ParaInfo) -> bool:
    if not FOOTNOTE_DEF_RE.match(p.text.strip()):
        return False
    # An inline superscript "(N)" marker is not a definition.
    return not _first_nonempty_superscript(p)


def parse_structure(path: Path) -> list[Page]:
    document = Document(str(path))
    pages: list[Page] = []
    current = Page(num=1)

    for paragraph in document.paragraphs:
        stripped = paragraph.text.strip()
        marker = PAGE_MARKER_RE.fullmatch(stripped)
        if marker:
            if current.raw_paras:
                pages.append(current)
            current = Page(num=int(marker.group(1)))
            continue
        current.raw_paras.append(_para_info(paragraph))

    if current.raw_paras or not pages:
        pages.append(current)
    return pages


def _split_body_and_footnotes(page: Page) -> None:
    in_footnotes = False
    current_fn: Optional[Footnote] = None
    for p in page.raw_paras:
        if not in_footnotes and _is_footnote_def(p):
            in_footnotes = True
        if not in_footnotes:
            page.body.append(p)
            continue
        m = FOOTNOTE_DEF_RE.match(p.text.strip())
        if m and not _first_nonempty_superscript(p):
            text = FOOTNOTE_DEF_RE.sub("", p.text.strip(), count=1).strip()
            current_fn = Footnote(page_num=page.num, local_num=int(m.group(1)), text=text)
            page.footnotes.append(current_fn)
        elif current_fn is not None:
            current_fn.text = (current_fn.text + " " + p.text.strip()).strip()
        else:
            page.body.append(p)


def _latin_letter_count(page: Page) -> int:
    count = 0
    for p in page.body:
        count += sum(1 for ch in p.text if ("A" <= ch <= "Z") or ("a" <= ch <= "z"))
    for fn in page.footnotes:
        count += sum(1 for ch in fn.text if ("A" <= ch <= "Z") or ("a" <= ch <= "z"))
    return count


@dataclass
class CleanStats:
    headers_dropped: int = 0
    cjk_stripped: int = 0


def _strip_cjk(text: str) -> tuple[str, int]:
    removed = len(CJK_RE.findall(text))
    return CJK_RE.sub("", text), removed


def _is_header_line(text: str) -> bool:
    """A short body line that is a running header, title, journal ref, or bare
    page number — OCR chrome to drop rather than content to keep."""
    t = text.strip()
    if not t:
        return False
    if BARE_PAGE_NUM_RE.match(t):
        return True
    if len(t.split()) > 14:
        return False
    low = t.lower()
    if any(kw in low for kw in HEADER_KEYWORDS):
        return True
    return bool(HEADER_MONTH_YEAR_RE.search(t))


def _clean_page_body(page: Page, stats: CleanStats) -> None:
    """Strip CJK OCR garbage and drop running-header lines, in place."""
    kept: list[ParaInfo] = []
    for p in page.body:
        clean_text, removed = _strip_cjk(p.text)
        stats.cjk_stripped += removed
        if _is_header_line(clean_text):
            stats.headers_dropped += 1
            continue
        # Mirror the CJK strip onto the runs used during reflow.
        for r in p.runs:
            r.text, _ = _strip_cjk(r.text)
        p.text = clean_text
        kept.append(p)
    page.body = kept
    for fn in page.footnotes:
        fn.text, removed = _strip_cjk(fn.text)
        stats.cjk_stripped += removed


# ---------------------------------------------------------------------------
# Step 2 / 3 — footnote mapping and body reflow
# ---------------------------------------------------------------------------


def _reflow_page(page: Page, local_to_global: dict[int, int]) -> str:
    out: list[str] = []
    found: set[int] = set()

    for p in page.body:
        for r in p.runs:
            txt = r.text
            if not txt:
                continue
            if r.highlighted or SEE_LOG_RE.search(txt):
                out.append(" <<A>> ")
                continue
            if r.superscript:
                num_match = re.fullmatch(r"\s*\(?(\d+)\)?\s*", txt)
                if num_match:
                    gid = local_to_global.get(int(num_match.group(1)))
                    if gid:
                        out.append(f" <<F{gid}>> ")
                        found.add(int(num_match.group(1)))
                        continue
            out.append(txt)
        out.append("\n")

    text = "".join(out)
    # Strip leaked tags and literal flag text; collapse the daggers to markers.
    text = SEE_LOG_RE.sub(" <<A>> ", text)
    text = LEAKED_TAG_RE.sub("", text)
    text = text.replace("†", " <<A>> ")

    # Fallback marker placement when no superscript markers were detected.
    if page.footnotes and not found:
        for fn in page.footnotes:
            m = re.search(r"\(%d\)" % fn.local_num, text)
            if m:
                text = f"{text[:m.start()]} <<F{fn.global_id}>> {text[m.end():]}"

    # Reflow: line breaks become spaces, collapse whitespace runs.
    text = re.sub(r"[ \t]*\n[ \t]*", " ", text)
    text = re.sub(r"\s+", " ", text).strip()
    # Final cleanup: drop any residual CJK, then un-glue sentence punctuation
    # that OCR line-joining left stuck to the next word.
    text = CJK_RE.sub("", text)
    text = GLUED_PUNCT_RE.sub(r"\1 \2", text)
    text = re.sub(r"\s{2,}", " ", text).strip()
    if text:
        text = f"<<P{page.num}>> {text}"
    return text


# ---------------------------------------------------------------------------
# Step 4 — reassemble without breaking sentences
# ---------------------------------------------------------------------------


def _ends_sentence(s: str) -> bool:
    t = re.sub(r"(<<[FPA]\d*>>|\s)+$", "", s.rstrip())
    t = t.rstrip("\"'”»)")
    return bool(t) and t[-1] in _SENTENCE_END


def _reassemble(pages: list[Page]) -> list[str]:
    paragraphs: list[str] = []
    buf = ""
    for page in pages:
        if not page.reflowed:
            continue
        if not buf:
            buf = page.reflowed
        elif _ends_sentence(buf):
            paragraphs.append(buf)
            buf = page.reflowed
        else:
            buf = f"{buf} {page.reflowed}"
    if buf:
        paragraphs.append(buf)
    return paragraphs


# ---------------------------------------------------------------------------
# Step 5 (fallback) — mechanical sentence-based paragraphing
# ---------------------------------------------------------------------------


def _split_sentences(text: str) -> list[str]:
    return [s.strip() for s in SENTENCE_SPLIT_RE.split(text) if s.strip()]


def _starts_discourse(sentence: str) -> bool:
    s = TOKEN_RE.sub("", sentence).lstrip(" \"“'‘(")
    return bool(DISCOURSE_RE.match(s))


def _split_long_paragraph(para: str) -> list[str]:
    if len(para) <= _PARA_MAX_CHARS:
        return [para]
    for m in re.finditer(r"[;:]", para):
        if m.start() >= _PARA_SPLIT_MIN_CHARS:
            head = para[: m.end()].strip()
            tail = para[m.end() :].strip()
            return [head] + _split_long_paragraph(tail) if tail else [head]
    return [para]


def _group_sentences(sentences: list[str]) -> list[str]:
    paragraphs: list[str] = []
    cur: list[str] = []
    cur_len = 0
    for s in sentences:
        # (a) Break BEFORE a discourse marker once the current paragraph has
        # content.
        if (
            cur
            and _starts_discourse(s)
            and len(cur) >= 1
            and cur_len >= _PARA_MARKER_MIN_CHARS
        ):
            paragraphs.append(" ".join(cur))
            cur, cur_len = [], 0
        cur.append(s)
        cur_len += len(s) + 1
        # (b) Break AFTER a boundary once the paragraph is long enough.
        if len(cur) >= _PARA_LEN_MIN_SENTENCES and cur_len >= _PARA_LEN_MIN_CHARS:
            paragraphs.append(" ".join(cur))
            cur, cur_len = [], 0
    if cur:
        paragraphs.append(" ".join(cur))

    # Last resort: hard-split any paragraph still over the ceiling.
    result: list[str] = []
    for para in paragraphs:
        result.extend(_split_long_paragraph(para))
    return result


def mechanical_paragraphs(step4_paragraphs: list[str]) -> list[str]:
    """Fallback paragraphing: reflow to one body, split to sentences, regroup.

    Never returns one-paragraph-per-page — it ignores the page-based grouping
    and rebuilds paragraphs from sentence boundaries and discourse markers.
    """
    body = " ".join(re.sub(r"\s+", " ", p).strip() for p in step4_paragraphs)
    body = GLUED_PUNCT_RE.sub(r"\1 \2", body)
    body = re.sub(r"\s{2,}", " ", body).strip()
    if not body:
        return []
    return _group_sentences(_split_sentences(body))


# ---------------------------------------------------------------------------
# Step 5 — intelligent paragraphing (Claude API)
# ---------------------------------------------------------------------------


def _normalize(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def _chunk_paragraphs(paragraphs: list[str], target_words: int = WORDS_PER_CHUNK):
    chunks: list[list[str]] = []
    current: list[str] = []
    words = 0
    for para in paragraphs:
        w = len(para.split())
        if current and words + w > target_words:
            chunks.append(current)
            current, words = [], 0
        current.append(para)
        words += w
    if current:
        chunks.append(current)
    return chunks


class FormatCheckpoint:
    """Checkpoint for the paragraphing step: ``{name}_format_checkpoint.json``."""

    def __init__(self, path: Path) -> None:
        self.path = path
        self.total: Optional[int] = None
        self.outputs: dict[int, str] = {}

    @classmethod
    def path_for(cls, input_path: Path) -> Path:
        return input_path.parent / f"{input_path.stem}_format_checkpoint.json"

    @classmethod
    def load(cls, input_path: Path) -> Optional["FormatCheckpoint"]:
        path = cls.path_for(input_path)
        if not path.exists():
            return None
        try:
            data = json.loads(path.read_text(encoding="utf-8"))
        except (json.JSONDecodeError, OSError):
            return None
        cp = cls(path)
        cp.total = data.get("total")
        cp.outputs = {int(k): v for k, v in data.get("outputs", {}).items()}
        return cp

    @classmethod
    def start(cls, input_path: Path) -> "FormatCheckpoint":
        return cls(cls.path_for(input_path))

    def _save(self) -> None:
        payload = {
            "total": self.total,
            "outputs": {str(k): v for k, v in self.outputs.items()},
        }
        self.path.write_text(
            json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    def ensure_total(self, total: int) -> None:
        if self.total is None:
            self.total = total
            self._save()

    def record(self, index: int, text: str) -> None:
        self.outputs[index] = text
        self._save()

    def delete(self) -> None:
        try:
            self.path.unlink()
        except FileNotFoundError:
            pass


def _paragraph_chunk(client, chunk_text: str) -> str:
    """One API call; returns the model's text (caller must verify it)."""
    response = client.messages.create(
        model=MODEL,
        max_tokens=8000,
        system=PARAGRAPHING_SYSTEM_PROMPT,
        messages=[{"role": "user", "content": chunk_text}],
    )
    return "".join(
        part.text for part in response.content if getattr(part, "type", "") == "text"
    )


def _call_with_retry(client, chunk_text: str) -> Optional[str]:
    try:
        return _paragraph_chunk(client, chunk_text)
    except Exception:  # noqa: BLE001 - retry once
        time.sleep(2)
        try:
            return _paragraph_chunk(client, chunk_text)
        except Exception:  # noqa: BLE001
            return None


def paragraph_body(
    paragraphs: list[str],
    client,
    checkpoint: FormatCheckpoint,
    progress=None,
) -> list[str]:
    """Return re-paragraphed text as a list of paragraphs.

    The model may only insert breaks; each chunk is verified and discarded if
    the model altered the text at all.
    """
    chunks = _chunk_paragraphs(paragraphs)
    checkpoint.ensure_total(len(chunks))
    processed: list[str] = []

    for i, chunk in enumerate(chunks):
        if progress is not None:
            progress(i + 1, len(chunks))
        chunk_text = "\n\n".join(chunk)

        if i in checkpoint.outputs:
            processed.append(checkpoint.outputs[i])
            continue

        returned = _call_with_retry(client, chunk_text)
        if returned is not None and _normalize(returned) == _normalize(chunk_text):
            result = returned.strip()
        else:
            # Reject: keep the chunk unbroken (its Step 4 paragraphing).
            result = chunk_text
        checkpoint.record(i, result)
        processed.append(result)

    combined = "\n\n".join(processed)
    return _split_into_paragraphs(combined)


def _split_into_paragraphs(text: str) -> list[str]:
    out: list[str] = []
    for block in re.split(r"\n\s*\n+", text):
        para = re.sub(r"\s*\n\s*", " ", block)
        para = re.sub(r"[ \t]+", " ", para).strip()
        if para:
            out.append(para)
    return out


# ---------------------------------------------------------------------------
# Step 6 — Markdown + Pandoc build
# ---------------------------------------------------------------------------

_MD_ESCAPE_RE = re.compile(r"([\\`*_{}\[\]()<>#+.!^~|-])")


def _escape_md(text: str) -> str:
    return _MD_ESCAPE_RE.sub(r"\\\1", text)


def _render_segment(seg: str) -> str:
    """Escape a token-free text segment and italicize Qur'anic {...} spans."""
    out: list[str] = []
    idx = 0
    for m in QURAN_RE.finditer(seg):
        out.append(_escape_md(seg[idx : m.start()]))
        out.append("*{" + _escape_md(m.group(1)) + "}*")
        idx = m.end()
    out.append(_escape_md(seg[idx:]))
    return "".join(out)


def _render_paragraph(text: str) -> str:
    parts: list[str] = []
    for piece in TOKEN_RE.split(text):
        if not piece:
            continue
        if piece.startswith("<<F"):
            parts.append(f"[^{piece[3:-2]}]")
        elif piece.startswith("<<P"):
            parts.append(f"^\\[{piece[3:-2]}\\]^")
        elif piece == "<<A>>":
            parts.append("^\\*^")
        else:
            parts.append(_render_segment(piece))
    return re.sub(r"\s{2,}", " ", "".join(parts)).strip()


def _derive_title(input_path: Path) -> str:
    stem = re.sub(r"_translated$", "", input_path.stem, flags=re.IGNORECASE)
    stem = stem.replace("_", " ").replace("-", " ")
    return re.sub(r"\s+", " ", stem).strip() or input_path.stem


def build_markdown(
    paragraphs: list[str],
    footnotes: list[Footnote],
    title: str,
) -> str:
    lines: list[str] = [f"# {title}", "", "*Translated from Arabic*", "", ""]
    for para in paragraphs:
        rendered = _render_paragraph(para)
        if rendered:
            lines.append(rendered)
            lines.append("")
    lines.append("")
    for fn in footnotes:
        note = _render_paragraph(fn.text) or "(no note text)"
        lines.append(f"[^{fn.global_id}]: {note}")
        lines.append("")
    return "\n".join(lines)


def _add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _set_doc_default_font(document, name: str, size_pt: int) -> None:
    """Force the document-default run font/size so nothing falls back to 10pt."""
    styles_el = document.styles.element
    path = "/".join((qn("w:docDefaults"), qn("w:rPrDefault"), qn("w:rPr")))
    rpr_default = styles_el.find(path)
    if rpr_default is None:
        return
    rfonts = rpr_default.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr_default.insert(0, rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs"):
        rfonts.set(qn(attr), name)
    half = str(int(size_pt * 2))
    for tag in ("w:sz", "w:szCs"):
        el = rpr_default.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            rpr_default.append(el)
        el.set(qn("w:val"), half)


def build_reference_doc(ref_path: Path) -> None:
    """Generate a Pandoc reference doc styled for a serif book/paper layout:
    US Letter, 1in margins, 12pt Times New Roman body (justified, 1.5 spacing,
    0.3in first-line indent, 8pt after), 10pt footnotes, centered page numbers.
    """
    with open(ref_path, "wb") as fh:
        subprocess.run(
            ["pandoc", "--print-default-data-file", "reference.docx"],
            stdout=fh,
            check=True,
        )
    document = Document(str(ref_path))

    # Document default (guards against any 10pt fallback).
    _set_doc_default_font(document, "Times New Roman", 12)

    # Body-text styles Pandoc may apply to paragraphs.
    for style_name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(12)
        pf = style.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf.first_line_indent = Inches(0.3)
        pf.line_spacing = 1.5
        pf.space_after = Pt(8)

    # Footnotes a touch smaller.
    for style_name in ("Footnote Text", "Footnote"):
        try:
            style = document.styles[style_name]
        except KeyError:
            continue
        style.font.name = "Times New Roman"
        style.font.size = Pt(10)

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        footer_para = section.footer.paragraphs[0]
        footer_para.text = ""
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_field(footer_para)

    document.save(str(ref_path))


def run_pandoc(md_path: Path, output_path: Path, ref_path: Path) -> None:
    subprocess.run(
        [
            "pandoc",
            str(md_path),
            "-f",
            "markdown",
            "-o",
            str(output_path),
            "--reference-doc",
            str(ref_path),
        ],
        check=True,
    )


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------


@dataclass
class FormatResult:
    output_path: Path
    pages_formatted: int
    footnotes_created: int
    asterisks: int
    skipped_pages: list[int]
    paragraphing_ran: bool
    headers_dropped: int = 0
    cjk_stripped: int = 0
    paragraphs_produced: int = 0
    median_paragraph_chars: int = 0


def pandoc_available() -> bool:
    return shutil.which("pandoc") is not None


def format_document(
    input_path: Path,
    client=None,
    restart: bool = False,
    progress=None,
) -> FormatResult:
    """Full Format-mode pipeline. Returns a :class:`FormatResult`.

    ``client`` is an Anthropic client or None (paragraphing skipped when None).
    Assumes Pandoc availability was already checked by the caller.
    """
    input_path = input_path.resolve()
    output_path = input_path.parent / f"{input_path.stem}_formatted.docx"

    if restart:
        FormatCheckpoint.path_for(input_path).unlink(missing_ok=True)

    # Step 1 — parse, clean (CJK + running headers), and classify pages.
    stats = CleanStats()
    pages = parse_structure(input_path)
    for page in pages:
        _split_body_and_footnotes(page)
        _clean_page_body(page, stats)
        if _latin_letter_count(page) < LATIN_LETTER_THRESHOLD:
            page.skipped = True

    skipped_pages = [p.num for p in pages if p.skipped]
    kept = [p for p in pages if not p.skipped]

    # Step 2 — assign global footnote ids across kept pages in order.
    footnotes: list[Footnote] = []
    next_id = 1
    page_local_maps: dict[int, dict[int, int]] = {}
    for page in kept:
        local_map: dict[int, int] = {}
        for fn in page.footnotes:
            fn.global_id = next_id
            next_id += 1
            local_map[fn.local_num] = fn.global_id
            footnotes.append(fn)
        page_local_maps[id(page)] = local_map

    # Step 3 — reflow each page into tokenised prose.
    for page in kept:
        page.reflowed = _reflow_page(page, page_local_maps[id(page)])

    # Step 4 — reassemble across page boundaries without breaking sentences.
    paragraphs = _reassemble(kept)

    # Step 5 — paragraphing. The Claude LLM pass is PREFERRED when a key is
    # present (it handles long run-on sentences); the mechanical, sentence-based
    # pass is the fallback and never produces one-paragraph-per-page.
    paragraphing_ran = False
    if client is not None and paragraphs:
        checkpoint = None if restart else FormatCheckpoint.load(input_path)
        if checkpoint is None:
            checkpoint = FormatCheckpoint.start(input_path)
        paragraphs = paragraph_body(paragraphs, client, checkpoint, progress=progress)
        checkpoint.delete()
        paragraphing_ran = True
    else:
        paragraphs = mechanical_paragraphs(paragraphs)

    asterisks = sum(p.count("<<A>>") for p in paragraphs)

    # Paragraph-length stats (visible characters, tokens removed).
    visible_lengths = [
        len(re.sub(r"\s{2,}", " ", TOKEN_RE.sub("", p)).strip()) for p in paragraphs
    ]
    median_len = int(statistics.median(visible_lengths)) if visible_lengths else 0

    # Step 6 — Markdown + Pandoc.
    title = _derive_title(input_path)
    markdown = build_markdown(paragraphs, footnotes, title)
    md_path = input_path.parent / f"{input_path.stem}_formatted.md"
    md_path.write_text(markdown, encoding="utf-8")
    ref_path = input_path.parent / f"{input_path.stem}_reference.docx"
    build_reference_doc(ref_path)
    run_pandoc(md_path, output_path, ref_path)

    # Clean up intermediates.
    md_path.unlink(missing_ok=True)
    ref_path.unlink(missing_ok=True)

    return FormatResult(
        output_path=output_path,
        pages_formatted=len(kept),
        footnotes_created=len(footnotes),
        asterisks=asterisks,
        skipped_pages=skipped_pages,
        paragraphing_ran=paragraphing_ran,
        headers_dropped=stats.headers_dropped,
        cjk_stripped=stats.cjk_stripped,
        paragraphs_produced=len(paragraphs),
        median_paragraph_chars=median_len,
    )
