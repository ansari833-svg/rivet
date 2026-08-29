"""Reconstruct output .docx files from translated blocks.

Two documents are produced:

* ``{name}_translated.docx`` - English-only, formatting preserved as closely as
  python-docx allows. Difficult flags are stripped and replaced with a subtle
  highlighted marker pointing the reader to the CSV log.
* ``{name}_bilingual.docx`` - Arabic/English parallel text with a title page,
  running headers, page numbers, real Word footnotes for every flagged or
  missing passage, and a "Translator's Notes" section at the end.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Optional
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .parser import Block, Element, ParsedDocument
from .translator import DIFFICULT_RE, extract_difficult_flags, strip_difficult_tags

# Arabic fonts that ship with Windows (primary + fallbacks).
ARABIC_FONT = "Traditional Arabic"
ARABIC_FONT_FALLBACKS = ("Arabic Typesetting", "Simplified Arabic")
ARABIC_SIZE_PT = 13
ENGLISH_FONT = "Times New Roman"
ENGLISH_SIZE_PT = 12

_FOOTNOTES_CT = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
)
_FOOTNOTES_RT = (
    "http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes"
)
_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------------------------------------------------------------------------
# Small OOXML helpers
# ---------------------------------------------------------------------------


def _set_rtl_paragraph(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    if p_pr.find(qn("w:bidi")) is None:
        p_pr.append(OxmlElement("w:bidi"))


def _set_arabic_run(run, size_pt: float = ARABIC_SIZE_PT) -> None:
    """Style a run as Arabic: RTL, complex-script font and size."""
    run.font.size = Pt(size_pt)
    run.font.name = ARABIC_FONT
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.append(r_fonts)
    r_fonts.set(qn("w:ascii"), ARABIC_FONT)
    r_fonts.set(qn("w:hAnsi"), ARABIC_FONT)
    r_fonts.set(qn("w:cs"), ARABIC_FONT)
    # Complex-script size.
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size_pt * 2)))
    r_pr.append(sz_cs)
    # Mark the run as right-to-left.
    rtl = OxmlElement("w:rtl")
    rtl.set(qn("w:val"), "1")
    r_pr.append(rtl)


def _add_horizontal_rule(paragraph) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAAAAA")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_page_number_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def _apply_paragraph_style(document, paragraph, block: Block) -> None:
    """Best-effort application of a heading level or named style."""
    if block.heading_level is not None:
        if block.heading_level == 0:
            style_name = "Title"
        else:
            style_name = f"Heading {min(block.heading_level, 9)}"
        try:
            paragraph.style = document.styles[style_name]
            return
        except KeyError:
            pass
    if block.style_name:
        try:
            paragraph.style = document.styles[block.style_name]
        except KeyError:
            pass


# ---------------------------------------------------------------------------
# Footnotes (real Word footnotes, added via low-level OOXML)
# ---------------------------------------------------------------------------


class FootnoteManager:
    """Accumulates footnotes and injects a footnotes part on :meth:`finalize`."""

    def __init__(self, document) -> None:
        self.document = document
        self._notes: list[tuple[int, str]] = []
        self._next_id = 1

    def add_footnote(self, paragraph, text: str) -> int:
        fid = self._next_id
        self._next_id += 1
        self._notes.append((fid, text))
        run = paragraph.add_run()
        r_pr = run._element.get_or_add_rPr()
        vert = OxmlElement("w:vertAlign")
        vert.set(qn("w:val"), "superscript")
        r_pr.append(vert)
        ref = OxmlElement("w:footnoteReference")
        ref.set(qn("w:id"), str(fid))
        run._element.append(ref)
        return fid

    def _footnotes_xml(self) -> bytes:
        parts = [
            '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
            f'<w:footnotes xmlns:w="{_W}">',
            '<w:footnote w:type="separator" w:id="-1">'
            "<w:p><w:r><w:separator/></w:r></w:p></w:footnote>",
            '<w:footnote w:type="continuationSeparator" w:id="0">'
            "<w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>",
        ]
        for fid, text in self._notes:
            body = escape(text)
            parts.append(
                f'<w:footnote w:id="{fid}">'
                '<w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
                "<w:r><w:rPr><w:rStyle w:val=\"FootnoteReference\"/>"
                '<w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
                f'<w:r><w:t xml:space="preserve"> {body}</w:t></w:r>'
                "</w:p></w:footnote>"
            )
        parts.append("</w:footnotes>")
        return "".join(parts).encode("utf-8")

    def finalize(self) -> None:
        if not self._notes:
            return
        partname = PackURI("/word/footnotes.xml")
        part = Part(
            partname,
            _FOOTNOTES_CT,
            self._footnotes_xml(),
            self.document.part.package,
        )
        self.document.part.relate_to(part, _FOOTNOTES_RT)


# ---------------------------------------------------------------------------
# Rendering helpers shared across the two output documents
# ---------------------------------------------------------------------------


def _is_missing(block_id: int, translations: dict[int, str], block: Block) -> bool:
    text = translations.get(block_id, "")
    return bool(block.text.strip()) and not text.strip()


def _selected_elements(
    parsed: ParsedDocument, selected_ids: set[int]
) -> list[Element]:
    out: list[Element] = []
    for el in parsed.elements:
        if el.kind == "para":
            if el.block_id in selected_ids:
                out.append(el)
        elif el.kind == "table" and el.grid is not None:
            if any(bid in selected_ids for row in el.grid for bid in row):
                out.append(el)
    return out


# ---------------------------------------------------------------------------
# English-only translated document
# ---------------------------------------------------------------------------


def _add_english_runs_with_markers(paragraph, translated: str) -> None:
    """Add the English text, replacing [DIFFICULT: ...] tags with a marker."""
    pos = 0
    for m in DIFFICULT_RE.finditer(translated):
        before = translated[pos : m.start()]
        if before:
            paragraph.add_run(before)
        marker = paragraph.add_run(" [†see log] ")
        marker.font.highlight_color = 7  # yellow (WD_COLOR_INDEX.YELLOW)
        marker.font.bold = True
        pos = m.end()
    tail = translated[pos:]
    if tail:
        paragraph.add_run(tail)


def build_translated_docx(
    parsed: ParsedDocument,
    translations: dict[int, str],
    selected_ids: set[int],
    output_path: Path,
) -> None:
    document = Document()
    elements = _selected_elements(parsed, selected_ids)

    for el in elements:
        if el.kind == "para" and el.block_id is not None:
            block = parsed.blocks[el.block_id]
            paragraph = document.add_paragraph()
            _apply_paragraph_style(document, paragraph, block)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT  # English is LTR

            if _is_missing(el.block_id, translations, block):
                run = paragraph.add_run(block.text)
                run.font.highlight_color = 7
            else:
                english = translations.get(el.block_id, block.text)
                _add_english_runs_with_markers(paragraph, english)
                # Approximate character formatting from the source block.
                for run in paragraph.runs:
                    if block.dominant_bold():
                        run.bold = True
                    if block.dominant_italic():
                        run.italic = True
                    if block.dominant_underline():
                        run.underline = True
                    size = block.first_size_pt()
                    if size:
                        run.font.size = Pt(size)
        elif el.kind == "table" and el.grid is not None:
            _build_table(document, parsed, translations, el.grid)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _build_table(document, parsed, translations, grid) -> None:
    rows = len(grid)
    cols = max((len(r) for r in grid), default=0)
    if rows == 0 or cols == 0:
        return
    table = document.add_table(rows=rows, cols=cols)
    try:
        table.style = "Table Grid"
    except KeyError:
        pass
    for r, row_ids in enumerate(grid):
        for c, bid in enumerate(row_ids):
            block = parsed.blocks[bid]
            cell = table.cell(r, c)
            if _is_missing(bid, translations, block):
                cell.text = block.text
            else:
                cell.text = strip_difficult_tags(translations.get(bid, block.text))


# ---------------------------------------------------------------------------
# Bilingual parallel-text document
# ---------------------------------------------------------------------------


@dataclass
class NoteItem:
    number: int
    original_arabic: str
    reason: str


def _configure_bilingual_sections(document, title: str) -> None:
    for section in document.sections:
        section.top_margin = Inches(1.25)
        section.bottom_margin = Inches(1.25)
        section.left_margin = Inches(1.25)
        section.right_margin = Inches(1.25)

        # Header: title (left) + label (right) via a right tab stop.
        header_para = section.header.paragraphs[0]
        header_para.text = ""
        text_width = section.page_width - section.left_margin - section.right_margin
        header_para.paragraph_format.tab_stops.add_tab_stop(
            text_width, WD_TAB_ALIGNMENT.RIGHT
        )
        run = header_para.add_run(f"{title}\tArabic–English Parallel Translation")
        run.font.size = Pt(9)
        run.italic = True

        # Footer: centered page number.
        footer_para = section.footer.paragraphs[0]
        footer_para.text = ""
        footer_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _add_page_number_field(footer_para)


def _add_title_page(document, filename: str, translation_date: str) -> None:
    for _ in range(6):
        document.add_paragraph()
    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(filename)
    run.bold = True
    run.font.size = Pt(24)

    for text, size in (
        ("Translated from Arabic", 14),
        (f"Date of translation: {translation_date}", 12),
        ("Translation assisted by Claude Sonnet (Anthropic)", 12),
    ):
        p = document.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.size = Pt(size)

    document.add_page_break()


def _add_arabic_block(document, block: Block) -> None:
    paragraph = document.add_paragraph()
    _apply_paragraph_style(document, paragraph, block)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    _set_rtl_paragraph(paragraph)
    run = paragraph.add_run(block.text)
    _set_arabic_run(run, ARABIC_SIZE_PT)


def _add_english_block_bilingual(
    document,
    block: Block,
    english: str,
    footnotes: FootnoteManager,
    notes_accumulator: list[NoteItem],
    missing: bool,
) -> None:
    paragraph = document.add_paragraph()
    _apply_paragraph_style(document, paragraph, block)
    if block.heading_level is None:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    def style_english(run) -> None:
        run.font.name = ENGLISH_FONT
        if block.heading_level is None:
            run.font.size = Pt(ENGLISH_SIZE_PT)

    if missing:
        run = paragraph.add_run(block.text)  # fallback to original Arabic
        _set_arabic_run(run, ARABIC_SIZE_PT)
        note_text = (
            "Translator's note: This passage was not translated in the "
            "automated run and requires manual review."
        )
        footnotes.add_footnote(paragraph, note_text)
        notes_accumulator.append(
            NoteItem(len(notes_accumulator) + 1, block.text, "Not returned by API")
        )
        return

    pos = 0
    for m in DIFFICULT_RE.finditer(english):
        before = english[pos : m.start()]
        if before:
            style_english(paragraph.add_run(before))
        original = m.group(1).strip()
        reason = m.group(2).strip()
        note_text = f"Translator's note: {original} — {reason}"
        footnotes.add_footnote(paragraph, note_text)
        notes_accumulator.append(
            NoteItem(len(notes_accumulator) + 1, original, reason)
        )
        pos = m.end()
    tail = english[pos:]
    if tail:
        style_english(paragraph.add_run(tail))
    if pos == 0 and not english.strip():
        style_english(paragraph.add_run(""))


def build_bilingual_docx(
    parsed: ParsedDocument,
    translations: dict[int, str],
    selected_ids: set[int],
    output_path: Path,
    title: str,
    translation_date: str,
) -> None:
    document = Document()
    _configure_bilingual_sections(document, title)
    _add_title_page(document, title, translation_date)

    footnotes = FootnoteManager(document)
    notes: list[NoteItem] = []

    elements = _selected_elements(parsed, selected_ids)
    for el in elements:
        block_ids: list[int] = []
        if el.kind == "para" and el.block_id is not None:
            block_ids = [el.block_id]
        elif el.kind == "table" and el.grid is not None:
            block_ids = [bid for row in el.grid for bid in row]

        for bid in block_ids:
            block = parsed.blocks[bid]
            if not block.text.strip():
                continue
            missing = _is_missing(bid, translations, block)
            english = translations.get(bid, block.text)

            _add_arabic_block(document, block)
            _add_english_block_bilingual(
                document, block, english, footnotes, notes, missing
            )
            _add_horizontal_rule(document.add_paragraph())

    _add_translators_notes_section(document, parsed, translations, selected_ids)

    footnotes.finalize()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(output_path))


def _add_translators_notes_section(
    document, parsed, translations, selected_ids
) -> None:
    document.add_page_break()
    heading = document.add_heading("Translator's Notes", level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    counter = 0
    for bid in sorted(selected_ids):
        block = parsed.blocks[bid]
        if _is_missing(bid, translations, block):
            counter += 1
            p = document.add_paragraph(style="List Number")
            p.add_run(
                "MISSING — not returned by the automated run; "
                "requires manual review. Original: "
            )
            r = p.add_run(block.text)
            _set_arabic_run(r, ARABIC_SIZE_PT)
            continue
        english = translations.get(bid, "")
        for flag in extract_difficult_flags(bid, english):
            counter += 1
            p = document.add_paragraph(style="List Number")
            r = p.add_run(flag.original_arabic)
            _set_arabic_run(r, ARABIC_SIZE_PT)
            p.add_run(f" — {flag.reason}")

    if counter == 0:
        document.add_paragraph("No sections were flagged for review.")


def output_paths(input_path: Path, range_suffix: str) -> dict[str, Path]:
    """Compute the standard output file paths for a given input + range."""
    stem = input_path.stem
    parent = input_path.parent
    return {
        "translated": parent / f"{stem}{range_suffix}_translated.docx",
        "bilingual_docx": parent / f"{stem}{range_suffix}_bilingual.docx",
        "bilingual_pdf": parent / f"{stem}{range_suffix}_bilingual.pdf",
        "log": parent / f"{stem}{range_suffix}_translation_log.csv",
    }
