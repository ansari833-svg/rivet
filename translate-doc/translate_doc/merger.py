"""Merge mode: combine several .docx and/or .txt files into one document.

Interactive flow (driven from :mod:`translate_doc.main`):

* scan a folder for .docx / .txt files (skipping ~$ Word lock files)
* natural-sort the list and let the user pick an explicit order
* confirm / correct the order
* determine the output format from the input types
* merge in order, inserting a subtle page break between documents
"""

from __future__ import annotations

from pathlib import Path

import natsort
from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn


def scan_folder(folder: Path) -> list[Path]:
    """Return sorted .docx/.txt files in ``folder``, skipping Word lock files."""
    files = [
        p
        for p in folder.iterdir()
        if p.is_file()
        and p.suffix.lower() in (".docx", ".txt")
        and not p.name.startswith("~$")
    ]
    # Natural sort so image-2 precedes image-10.
    return natsort.natsorted(files, key=lambda p: p.name)


def determine_output_format(files: list[Path]) -> str:
    """Return 'docx' or 'rtf' based on the selected input types."""
    suffixes = {p.suffix.lower() for p in files}
    if suffixes == {".docx"}:
        return "docx"
    return "rtf"


def is_mixed_types(files: list[Path]) -> bool:
    suffixes = {p.suffix.lower() for p in files}
    return ".docx" in suffixes and ".txt" in suffixes


def has_translated_and_bilingual(files: list[Path]) -> bool:
    names = [p.stem.lower() for p in files]
    has_translated = any("_translated" in n for n in names)
    has_bilingual = any("_bilingual" in n for n in names)
    return has_translated and has_bilingual


# ---------------------------------------------------------------------------
# .docx merging
# ---------------------------------------------------------------------------


def _append_docx_body(target: Document, source_path: Path, add_break: bool) -> None:
    source = Document(str(source_path))
    if add_break:
        target.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    for element in source.element.body:
        # Skip the trailing sectPr so we don't inherit the source's section.
        if element.tag == qn("w:sectPr"):
            continue
        target.element.body.append(element)


def merge_docx(files: list[Path], output_path: Path) -> None:
    """Merge .docx files by appending body XML (never raw bytes)."""
    target = Document(str(files[0]))
    for source_path in files[1:]:
        _append_docx_body(target, source_path, add_break=True)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    target.save(str(output_path))


# ---------------------------------------------------------------------------
# .txt -> .rtf merging
# ---------------------------------------------------------------------------


def _rtf_escape(text: str) -> str:
    """Escape RTF control chars and emit non-ASCII as \\uNNNN? escapes."""
    out: list[str] = []
    for ch in text:
        if ch in ("\\", "{", "}"):
            out.append("\\" + ch)
        elif ch == "\n":
            out.append("\\par\n")
        elif ch == "\r":
            continue
        elif ch == "\t":
            out.append("\\tab ")
        elif ord(ch) < 128:
            out.append(ch)
        else:
            code = ord(ch)
            if code > 32767:
                code -= 65536  # RTF uses signed 16-bit
            out.append(f"\\u{code}?")
    return "".join(out)


def merge_txt_to_rtf(files: list[Path], output_path: Path) -> None:
    """Merge .txt files into a single valid RTF document."""
    body_parts: list[str] = []
    for i, path in enumerate(files):
        text = path.read_text(encoding="utf-8", errors="replace")
        if i > 0:
            body_parts.append("\\page\n")  # subtle page break between documents
        body_parts.append(_rtf_escape(text))
        body_parts.append("\\par\n")

    rtf = (
        "{\\rtf1\\ansi\\ansicpg1252\\deff0"
        "{\\fonttbl{\\f0 Times New Roman;}}\n"
        "\\f0\\fs24\n" + "".join(body_parts) + "}"
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    # RTF control words are ASCII; \\uNNNN? escapes carry the Unicode payload.
    output_path.write_text(rtf, encoding="ascii", errors="ignore")


def merge(files: list[Path], output_path: Path) -> str:
    """Merge according to the detected format; returns the format used."""
    fmt = determine_output_format(files)
    if fmt == "docx":
        merge_docx(files, output_path)
    else:
        merge_txt_to_rtf(files, output_path)
    return fmt
